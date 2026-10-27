# -*- coding: utf-8 -*-
"""Fase 2 langkah 7: filling a large table of a page from a spreadsheet (platform/admin/impor_tabel.py, K-P34).

- which tables qualify (shown as a table, at least 50 rows, no required file column) and that every column name of
  each is unambiguous;
- the template: its sheets, the rows the page holds now, drop-down lists, files left out, a text starting with "="
  kept as text, a control character dropped;
- reading: a template filled and sent back gives the same rows; headers in any order and case, unknown columns
  ignored, every kind of cell (numbers, years, Excel percentages, choices without case, ya/tidak, dates in text
  columns, line breaks in one-line texts), the page's checks per row (required, spam, https, limits) with the row
  number of the file, CSV; files that cannot be used at all, with the reason;
- the rows after an import: added without rows already there, or replacing them while an unchanged row keeps its PDF;
  the preview's numbers; the table's limit;
- the admin: the link in the editor only where a table qualifies, the import page, the template with the draft's rows,
  a preview that stores nothing, the import into the draft (published content unchanged, history kept, audit log),
  a preview used once and only by its own account, replacing, the limit, refused files.

Needs the packages of platform/admin/requirements.txt: run with .venv/bin/python.
"""
from __future__ import print_function

import datetime
import io
import json
import os
import re
import shutil
import sys
import tempfile

AKAR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for sub in ("platform", "platform/admin", "platform/bangun", "platform/pemantau"):
    sys.path.insert(0, os.path.join(AKAR, *sub.split("/")))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import flask  # noqa: F401
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("uji_impor_tabel: paket admin belum terpasang; jalankan dengan .venv/bin/python (lihat README)")
    sys.exit(1)

KERJA = tempfile.mkdtemp(prefix="prodi-uji-impor-tabel-")
for nama, sub in (("PRODI_DATA", "data"), ("PRODI_RIWAYAT", "riwayat"), ("PRODI_RAHASIA", "rahasia"), ("PRODI_LOG", "log"),
                  ("WWW_DATA", "www"), ("PRODI_PUBLIK", "publik"), ("PRODI_NGINX", "nginx"), ("PRODI_UNGGAHAN", "unggahan")):
    os.environ[nama] = os.path.join(KERJA, sub)
os.environ["PRODI_MANIFES"] = os.path.join(KERJA, "manifes.json")
os.environ["PRODI_PENYIMPANAN"] = "lokal"
os.environ.pop("PRODI_MODE", None)

import akun  # noqa: E402
import app as app_mod  # noqa: E402
import impor_tabel as it  # noqa: E402
import isi_halaman as ih  # noqa: E402
import media  # noqa: E402
import registri  # noqa: E402
import riwayat  # noqa: E402
import tarik_www  # noqa: E402
from formulir_uji import pdf, prodi  # noqa: E402

A, B = "55201", "20201"
gagal = []


def cek(kondisi, pesan):
    if not kondisi:
        gagal.append(pesan)


def tabel(id_, kunci):
    return dict(it.tabel_halaman(id_))[kunci]


def xlsx(baris, format_sel=None):
    """An Excel file of rows (the first is the header); format_sel: {(row, column): number format}, 1-based."""
    buku = Workbook()
    ws = buku.active
    for b in baris:
        ws.append(b)
    for (r, c), f in (format_sel or {}).items():
        ws.cell(row=r, column=c).number_format = f
    keluar = io.BytesIO()
    buku.save(keluar)
    return keluar.getvalue()


def galat_nilai(fungsi):
    try:
        fungsi()
    except ValueError as e:
        return str(e)
    return None


try:
    # ============================================================== which tables
    semua = {(id_, k) for id_ in ih.TEMPLAT for k, _n in it.tabel_halaman(id_)}
    janji = {("kurikulum", "mata_kuliah"), ("kerja-sama", "daftar"), ("penelitian", "daftar"), ("publikasi", "daftar"),
             ("pkm", "daftar"), ("luaran", "daftar"), ("organisasi-prestasi", "prestasi"),
             ("pengembangan-sdm", "pelatihan"), ("pengembangan-sdm", "prestasi"), ("survei-kepuasan", "hasil"),
             ("jadwal", "jadwal"), ("tendik", "daftar")}
    cek(janji <= semua, "tabel yang dijanjikan bisa diimpor: kurang %r" % sorted(janji - semua))
    cek(("visi-misi", "dokumen") not in semua and ("laboratorium", "daftar") not in semua
        and ("organisasi-prestasi", "organisasi") not in semua and ("alumni", "tracer") not in semua
        and not any(k == "tambahan" for _i, k in semua), "tanpa tabel dokumen, kartu, tabel kecil, bagian tambahan")
    for id_, k in sorted(semua):
        node = tabel(id_, k)
        cek(node["tampil"] == "tabel" and node["maks"] >= it.MIN_BARIS and not any(
            n["jenis"] in it.MEDIA and n.get("wajib") for _k, n in node["isi"]["bidang"]), "%s %s memenuhi syarat" % (id_, k))
        cek(all(n["jenis"] not in it.MEDIA for _k, n in it.kolom(node)), "%s %s: kolom berkas tidak ada di Excel" % (id_, k))
        nama = {}
        for kk, n in it.kolom(node):
            for x in (kk, n["label"]):
                nama.setdefault(it._nama_kolom(x), set()).add(kk)
        cek(all(len(v) == 1 for v in nama.values()), "%s %s: nama kolom tidak ganda: %r" % (id_, k, nama))
    cek(it.tabel_halaman("tidak-ada") == [] and it.tabel_halaman("visi-misi") == [], "halaman tanpa tabel")
    ih.TEMPLAT["uji-foto"] = ih.halaman(("foto_wajib", ih.tabel("Foto", [("nama", ih.teks("Nama")), ("foto", ih.foto())], 100)),
                                        ("foto_boleh", ih.tabel("Foto", [("nama", ih.teks("Nama")), ("foto", ih.foto(wajib=False))], 100)))
    cek([k for k, _n in it.tabel_halaman("uji-foto")] == ["foto_boleh"], "tabel dengan kolom berkas wajib tidak bisa diimpor")
    del ih.TEMPLAT["uji-foto"]

    # ============================================================== the template
    mk = tabel("kurikulum", "mata_kuliah")
    ada = [{"semester": 1, "kode": "=HYPERLINK(\"https://contoh.test\")", "nama": "Algoritma\x07", "sks": 3, "sifat": "Wajib",
            "rps": {"jenis": "pdf", "jalur": "%s/202610-0123456789ab.pdf" % A}},
           {"semester": 2, "nama": "Basis Data", "sks": 3}]
    buku = load_workbook(io.BytesIO(it.templat_xlsx(mk, ada, "Mata kuliah — Kurikulum")))
    cek(buku.sheetnames == ["Data", "Petunjuk", "Pilihan"] and buku["Pilihan"].sheet_state == "hidden",
        "templat: lembar Data, Petunjuk, Pilihan (tersembunyi)")
    data = [[c.value for c in b] for b in buku["Data"].iter_rows()]
    cek(data[0] == ["Semester", "Kode", "Mata kuliah", "SKS", "Sifat"], "templat: kepala kolom tanpa RPS: %r" % data[0])
    cek(data[1] == [1, '=HYPERLINK("https://contoh.test")', "Algoritma", 3, "Wajib"] and data[2][:4] == [2, None, "Basis Data", 3],
        "templat: baris yang ada, karakter kendali dibuang: %r" % data[1:])
    cek(buku["Data"]["B2"].data_type == "s", "templat: teks yang diawali = tetap teks, bukan rumus")
    dv = buku["Data"].data_validations.dataValidation
    cek(len(dv) == 1 and str(dv[0].sqref) == "E2:E251" and "Pilihan" in dv[0].formula1, "templat: daftar pilihan Sifat")
    petunjuk = "\n".join(" | ".join(str(v) for v in b if v is not None) for b in buku["Petunjuk"].iter_rows(values_only=True))
    cek("Kolom RPS tidak ada di Excel" in petunjuk and "Semester | ya | Angka bulat 1 sampai 14" in petunjuk
        and "Sifat | Salah satu: Wajib; Pilihan" in petunjuk and "paling banyak 250 baris" in petunjuk, "templat: petunjuk")
    pn = tabel("penelitian", "daftar")
    dv = load_workbook(io.BytesIO(it.templat_xlsx(pn, [{"judul": "J", "mahasiswa": True}], "P")))["Data"]
    cek(dv["E2"].value == "ya" and len(dv.data_validations.dataValidation) == 1, "templat: centang sebagai ya, dengan daftar ya/tidak")

    # ============================================================== reading: a template sent back
    sv = tabel("survei-kepuasan", "hasil")
    baris_sv = [{"survei": "Mahasiswa", "tahun": 2025, "aspek": "Layanan akademik", "sangat_baik": 40, "baik": 45.5,
                 "cukup": 10, "kurang": 4.5}, {"survei": "Pengguna lulusan", "tahun": 2024, "aspek": "Etika", "baik": 100}]
    hasil = it.rencana(sv, it.baca_berkas("survei.xlsx", it.templat_xlsx(sv, baris_sv, "S")), A)
    cek(hasil["baris"] == baris_sv and not hasil["galat"] and hasil["dibaca"] == 2, "templat bolak-balik: %r" % hasil)
    pb = tabel("publikasi", "daftar")
    baris_pb = [{"judul": "Deteksi dini", "penulis": "A. Uji; B. Coba", "wadah": "Jurnal Uji", "tahun": 2025,
                 "tautan": "https://doi.org/10.1234/uji"}]
    cek(it.rencana(pb, it.baca_berkas("p.xlsx", it.templat_xlsx(pb, baris_pb, "P")), A)["baris"] == baris_pb,
        "templat bolak-balik: tautan")
    hasil = it.rencana(pn, it.baca_berkas("p.xlsx", it.templat_xlsx(pn, [{"judul": "J", "tahun": 2020, "mahasiswa": True}], "P")), A)
    cek(hasil["baris"] == [{"judul": "J", "tahun": 2020, "mahasiswa": True}], "templat bolak-balik: centang: %r" % hasil)

    # ============================================================== reading: the cells of a file made elsewhere
    berkas = xlsx([
        ["TAHUN", "Survei", "aspek", "Catatan", "Sangat baik", "baik", "Cukup", "kurang"],
        [2024.0, "mahasiswa", "Layanan\nakademik", "abaikan", 0.875, "10,5%", None, "2"],
        [],
        ["2.025", "Alumni", datetime.date(2025, 9, 1), "", "87,5", None, None, None],
        ["dua ribu", "Siapa saja", "Sarana", "", 101, "x", None, None],
        [2025, "Dosen", "Promo slot gacor", "", None, None, None, None],
        [2025, "Dosen", "", "", None, None, None, None],
        [None, None, None, "hanya catatan", None, None, None, None]], {(2, 5): "0.0%"})
    hasil = it.rencana(sv, it.baca_berkas("survei.xlsx", berkas), A)
    cek(hasil["baris"] == [{"tahun": 2024, "survei": "Mahasiswa", "aspek": "Layanan akademik", "sangat_baik": 87.5,
                            "baik": 10.5, "kurang": 2},
                           {"tahun": 2025, "survei": "Alumni", "aspek": "01/09/2025", "sangat_baik": 87.5}],
        "sel dibaca seperti formulir: %r" % hasil["baris"])
    cek(hasil["abaikan"] == ["Catatan"] and hasil["dibaca"] == 5 and hasil["bermasalah"] == 3,
        "kolom lain diabaikan, baris yang hanya berisi kolom lain dilewati: %r" % {k: hasil[k] for k in ("abaikan", "dibaca", "bermasalah")})
    g = hasil["galat"]
    for harus in ("Baris 5, Tahun: harus angka bulat", "Baris 5, Responden: harus salah satu dari: Mahasiswa;",
                  "Baris 5, Sangat baik (%): harus angka 0 sampai 100", "Baris 5, Baik (%): harus angka 0 sampai 100",
                  "Baris 6, Aspek: memuat teks judi", "Baris 7, Aspek: wajib diisi"):
        cek(any(p.startswith(harus) for p in g), "pesan %r: %r" % (harus, g))
    cek(not any(p.startswith("Baris 5, Tahun: wajib") for p in g), "satu pesan per kolom: %r" % g)
    pesan = galat_nilai(lambda: it.rencana(pb, it.baca_berkas("p.xlsx", xlsx([["Judul", "Tahun"], ["X", 2025]])), A))
    cek(pesan is None, "berkas tanpa kolom opsional diterima: %r" % pesan)
    hasil = it.rencana(pb, it.baca_berkas("p.xlsx", xlsx([["Judul", "Tautan", "Tahun"], ["X", "www.contoh.ac.id", 2025]])), A)
    cek(hasil["galat"] == ["Baris 2, Tautan (DOI/URL): harus alamat lengkap yang diawali https://"], "tautan tanpa https: %r" % hasil["galat"])
    hasil = it.rencana(pn, it.baca_berkas("p.xlsx", xlsx([["Judul", "Tahun", "Melibatkan mahasiswa"], ["A", 2025, "Ya"],
                                                          ["B", 2025, "x"], ["C", 2025, ""], ["D", 2025, "tidak"],
                                                          ["E", 2025, "mungkin"], ["F", 2025, True]])), A)
    cek([r.get("mahasiswa") for r in hasil["baris"]] == [True, True, None, None, True]
        and hasil["galat"] == ["Baris 6, Melibatkan mahasiswa: isi ya atau tidak"], "ya/tidak: %r" % hasil)
    jd = tabel("jadwal", "jadwal")
    hasil = it.rencana(jd, it.baca_berkas("j.xlsx", xlsx([["Kegiatan", "Waktu", "Tempat"],
                                                          ["Ujian", datetime.time(8, 30), "R1"],
                                                          ["Wisuda", datetime.datetime(2025, 10, 4, 9, 0), "Auditorium"]])), A)
    cek([r.get("waktu") for r in hasil["baris"]] == ["08:30", "04/10/2025 09:00"], "jam dan tanggal berjam di kolom teks: %r" % hasil)
    csv_ = "\ufeffSemester;Kode;Mata kuliah;SKS;Sifat\n1;IF101;Pengantar Teknologi Informasi;2;wajib\n".encode("utf-8")
    cek(it.rencana(mk, it.baca_berkas("mk.csv", csv_), A)["baris"] == [
        {"semester": 1, "kode": "IF101", "nama": "Pengantar Teknologi Informasi", "sks": 2, "sifat": "Wajib"}], "CSV dengan titik koma")

    # files that cannot be used at all
    for harus, fungsi in (
            ("baris pertama berkas harus berisi nama kolom", lambda: it.rencana(mk, it.baca_berkas("x.xlsx", xlsx([["a", "b"], [1, 2]])), A)),
            ("kolom wajib tidak ada: Mata kuliah, SKS", lambda: it.rencana(mk, it.baca_berkas("x.xlsx", xlsx([["Semester"], [1]])), A)),
            ("paling banyak 250 baris", lambda: it.rencana(mk, it.baca_berkas("x.xlsx", xlsx(
                [["Semester", "Mata kuliah", "SKS"]] + [[1, "MK %d" % i, 2] for i in range(251)])), A)),
            ("Format .xls lama", lambda: it.baca_berkas("lama.xls", b"\xd0\xcf\x11\xe0" + b"0" * 100)),
            ("Berkas harus Excel", lambda: it.baca_berkas("x.pdf", pdf())),
            ("Berkas Excel rusak", lambda: it.baca_berkas("x.xlsx", b"PK\x03\x04bukan zip")),
            ("Berkas terlalu besar", lambda: it.baca_berkas("x.csv", b"a" * (2 * 1024 * 1024 + 1))),
            ("Berkas kosong", lambda: it.baca_berkas("x.xlsx", xlsx([[None], []]))),
            ("Pilih berkas", lambda: it.baca_berkas("x.xlsx", b""))):
        pesan = galat_nilai(fungsi)
        cek(pesan is not None and harus in pesan, "berkas ditolak dengan alasan %r: %r" % (harus, pesan))

    # ============================================================== the rows after an import
    rps = {"jenis": "pdf", "jalur": "%s/202610-0123456789ab.pdf" % A, "judul": "RPS", "halaman": 1, "ukuran": 10}
    lama = [{"semester": 1, "kode": "IF101", "nama": "Algoritma  Dasar", "sks": 3, "sifat": "", "rps": rps},
            {"semester": 2, "nama": "Basis Data", "sks": 3, "rps": rps}]
    baru = [{"semester": 1, "kode": "IF101", "nama": "Algoritma Dasar", "sks": 3}, {"semester": 3, "nama": "Jaringan", "sks": 3}]
    cek(it.ringkasan(mk, lama, baru) == {"lama": 2, "sama": 1, "tambah": 3, "ganti": 2, "lepas": 1, "maks": 250},
        "ringkasan pratinjau: %r" % it.ringkasan(mk, lama, baru))
    cek(it.gabung(mk, lama, baru, "tambah") == lama + [baru[1]], "tambah: baris yang sudah ada tidak dua kali")
    cek(it.gabung(mk, lama, baru, "ganti") == [dict(baru[0], rps=rps), baru[1]], "ganti: baris yang sama tetap membawa RPS")
    ru = tabel("ruang", "daftar")
    penuh = [{"jenis": "Ruang %d" % i} for i in range(30)]
    pesan = galat_nilai(lambda: it.gabung(ru, penuh, [{"jenis": "Lain %d" % i} for i in range(21)], "tambah"))
    cek(pesan and "paling banyak 50 baris" in pesan, "batas tabel saat menambah: %r" % pesan)
    cek(galat_nilai(lambda: it.gabung(ru, penuh, [], "hapus")) is not None, "pilihan impor yang tidak dikenal ditolak")

    # ============================================================== the admin
    os.makedirs(os.environ["WWW_DATA"])
    for nama, isi in {"prodi": {"program": [prodi(A, "Teknik Informatika"), prodi(B, "Teknik Elektro")]},
                      "unit": {"unit": [{"kode": "ft", "nama": "Fakultas Teknik", "jenis": "fakultas",
                                         "bendera": {"dasar": "#000099", "tulisan": "#FFFF00"}}]},
                      "kalender": {"semester": []}, "situs": {"identitas": {}, "kontak": {}}}.items():
        with io.open(os.path.join(os.environ["WWW_DATA"], nama + ".json"), "w", encoding="utf-8") as f:
            json.dump(isi, f)
    tarik_www.terapkan("penguji")
    path_csv = os.path.join(KERJA, "slug.csv")
    with io.open(path_csv, "w", encoding="utf-8") as f:
        f.write(",".join(registri.slug_mod.KOLOM) + "\n%s,S1,Teknik Informatika,FT,s1informatika,x,x\n"
                "%s,S1,Teknik Elektro,FT,s1elektro,x,x\n" % (A, B))
    registri.siapkan_dari_csv("penguji", path_csv)
    SANDI = {}

    def akun_baru(nama, peran, prodi_, pelaku=None):
        id_a, sementara = akun.buat(pelaku, nama, nama.capitalize() + " Uji", peran, prodi_)
        SANDI[nama] = "Kebun 7 pagi hari nomor %d" % (len(SANDI) + 1)
        akun.ganti_sandi(id_a, sementara, SANDI[nama])
        return akun.ambil(id_a)
    super_ = akun_baru("superuji", "super", [])
    for nama, peran, p in (("adma", "admin", [A]), ("opa", "operator", [A]), ("opb", "operator", [B]),
                           ("opab", "operator", [A, B])):
        akun_baru(nama, peran, p, super_)
    app = app_mod.buat_app({"SECRET_KEY": "k" * 64, "ADMIN_IP_BOLEH": "103.84.116.0/24", "PRODI_BANGUN": "mati"})
    klien = {}
    for i, nama in enumerate(("adma", "opa", "opb", "opab")):
        c = app.test_client()
        c.environ_base["REMOTE_ADDR"] = "103.84.116.%d" % (70 + i)
        c.get("/masuk")
        with c.session_transaction() as s:
            token = s["csrf"]
        cek(c.post("/masuk", data={"csrf": token, "pengguna": nama, "sandi": SANDI[nama]}).status_code == 302, "%s masuk" % nama)
        klien[nama] = c
    ka, ko, kb, kab = klien["adma"], klien["opa"], klien["opb"], klien["opab"]

    def kirim(c, url, **data):
        with c.session_transaction() as s:
            data["csrf"] = s["csrf"]
        return c.post(url, data=data, content_type="multipart/form-data")

    def unggah(c, id_, kunci, isi, nama="isi.xlsx", kode=A):
        return kirim(c, "/prodi/%s/halaman/%s/impor" % (kode, id_), tabel=kunci, berkas=(io.BytesIO(isi), nama))

    def token_dari(r):
        m = re.search(r'name="token" value="([A-Za-z0-9_-]+)"', r.get_data(as_text=True))
        return m.group(1) if m else None

    def terapkan(c, id_, token, mode, kode=A):
        return kirim(c, "/prodi/%s/halaman/%s/impor/terapkan" % (kode, id_), token=token or "", mode=mode)

    URL = "/prodi/%s/halaman/kurikulum" % A
    cek('href="%s/impor">Impor dari Excel</a>' % URL in ko.get(URL).get_data(as_text=True)
        and "Tabel Mata kuliah bisa diisi dari Excel" in ko.get(URL).get_data(as_text=True), "penyunting: tautan impor")
    cek("Impor dari Excel" not in ko.get("/prodi/%s/halaman/visi-misi" % A).get_data(as_text=True), "tanpa tautan di halaman tanpa tabel besar")
    cek(ko.get("/prodi/%s/halaman/visi-misi/impor" % A).status_code == 404, "halaman tanpa tabel besar: 404")
    cek(ko.get(URL + "/impor/dokumen.xlsx").status_code == 404, "tabel yang tidak bisa diimpor: 404")
    cek(kb.get(URL + "/impor").status_code == 403, "operator prodi lain: 403")

    rps_a = media.simpan_pdf(A, pdf(), "RPS Algoritma")
    awal = [{"semester": 1, "kode": "IF101", "nama": "Algoritma", "sks": 3, "sifat": "Wajib", "rps": rps_a}]
    ih.simpan_draf(A, "kurikulum", {"struktur": "Kurikulum 2025.", "mata_kuliah": awal}, "opa")
    teks = ko.get(URL + "/impor").get_data(as_text=True)
    cek("Sekarang 1 baris; paling banyak 250 baris" in teks and "Kolom RPS tidak ada di" in teks
        and 'href="%s/impor/mata_kuliah.xlsx"' % URL in teks and "Baca berkas" in teks, "halaman impor")
    r = ko.get(URL + "/impor/mata_kuliah.xlsx")
    cek(r.status_code == 200 and r.headers["Content-Type"] == it.TIPE_XLSX
        and 'filename="%s-kurikulum-mata_kuliah.xlsx"' % A in r.headers["Content-Disposition"], "unduhan templat")
    buku = load_workbook(io.BytesIO(r.data))
    cek([c.value for c in buku["Data"][2]] == [1, "IF101", "Algoritma", 3, "Wajib"], "templat berisi baris draf")

    # the template filled: one row back unchanged, two new, one with a problem
    ws = buku["Data"]
    for b in ([1, "IF102", "Kalkulus", 3, "Wajib"], [2, "IF201", "Struktur Data", 3, "wajib"], [3, None, None, 2, "Pilihan"]):
        ws.append(b)
    keluar = io.BytesIO()
    buku.save(keluar)
    versi = ih.baca(A, "kurikulum")["versi"]
    r = unggah(ko, "kurikulum", "mata_kuliah", keluar.getvalue())
    teks = r.get_data(as_text=True)
    cek(r.status_code == 200 and "4 baris dibaca, <strong>3</strong> siap diimpor" in teks
        and "Baris 5, Mata kuliah: wajib diisi" in teks and "Tabel menjadi 3 baris. 1 baris di berkas sama persis" in teks
        and "Belum ada yang disimpan" in teks, "pratinjau: %s" % re.sub(r"\s+", " ", teks)[-1500:])
    cek(ih.baca(A, "kurikulum")["versi"] == versi, "pratinjau tidak menyimpan apa pun")
    token = token_dari(r)
    cek(terapkan(ka, "kurikulum", token, "tambah").status_code == 302 and ih.baca(A, "kurikulum")["versi"] == versi,
        "pratinjau milik akun lain tidak berlaku")
    r = unggah(ko, "kurikulum", "mata_kuliah", keluar.getvalue())
    token = token_dari(r)
    r = terapkan(ko, "kurikulum", token, "tambah")
    data = ih.baca(A, "kurikulum")
    mk_draf = ih.draf(data)["isi"]["mata_kuliah"]
    cek(r.status_code == 302 and r.headers["Location"].endswith(URL) and [b["nama"] for b in mk_draf] ==
        ["Algoritma", "Kalkulus", "Struktur Data"] and mk_draf[0]["rps"] == rps_a and mk_draf[2]["sifat"] == "Wajib",
        "tambah: baris baru di bawah, baris lama tetap dengan RPS: %r" % mk_draf)
    cek(ih.draf(data)["isi"]["struktur"] == "Kurikulum 2025." and data.get("terbit") is None and ih.draf(data)["oleh"] == "opa",
        "impor hanya mengubah tabelnya, di draf; situs tidak berubah")
    teks = ko.get(URL).get_data(as_text=True)
    cek("Mata kuliah: 2 baris ditambahkan (1 baris yang sudah ada dilewati); tabel kini 3 baris" in teks, "pesan sesudah impor")
    cek(any(v["berkas"] == "prodi/%s/halaman-kurikulum" % A for v in riwayat.daftar(500)), "versi sebelum impor ada di riwayat")
    r = terapkan(ko, "kurikulum", token, "tambah")
    cek(r.status_code == 302 and "tidak berlaku" in ko.get(r.headers["Location"]).get_data(as_text=True)
        and len(ih.draf(ih.baca(A, "kurikulum"))["isi"]["mata_kuliah"]) == 3, "pratinjau hanya sekali pakai")
    cek("isi-tabel-diimpor" in ka.get("/prodi/%s/audit" % A).get_data(as_text=True), "impor tercatat di log audit")

    # replacing: an unchanged row keeps its PDF
    r = unggah(ka, "kurikulum", "mata_kuliah", xlsx([["Semester", "Kode", "Mata kuliah", "SKS", "Sifat"],
                                                      [1, "IF101", "Algoritma", 3, "Wajib"], [4, "IF401", "Skripsi", 6, "Wajib"]]))
    teks = r.get_data(as_text=True)
    cek("Tabel menjadi 2 baris; 3 baris yang ada sekarang diganti." in teks and "tetap membawa berkasnya (RPS)" in teks,
        "pratinjau ganti")
    terapkan(ka, "kurikulum", token_dari(r), "ganti")
    mk_draf = ih.draf(ih.baca(A, "kurikulum"))["isi"]["mata_kuliah"]
    cek([b["nama"] for b in mk_draf] == ["Algoritma", "Skripsi"] and mk_draf[0]["rps"] == rps_a and "rps" not in mk_draf[1],
        "ganti semua: %r" % mk_draf)

    # a preview belongs to its programme and its page
    isi_pb = xlsx([["Judul", "Tahun"], ["Publikasi uji", 2025]])
    for kode_lain, id_lain in ((B, "publikasi"), (A, "penelitian")):
        r = terapkan(kab, id_lain, token_dari(unggah(kab, "publikasi", "daftar", isi_pb)), "tambah", kode_lain)
        cek(r.status_code == 302 and ih.draf(ih.baca(kode_lain, id_lain)) is None,
            "pratinjau %s publikasi tidak berlaku di %s %s" % (A, kode_lain, id_lain))
    cek(ih.draf(ih.baca(A, "publikasi")) is None, "pratinjau yang dipakai di tempat lain hilang")

    # the limit of a table, and files refused
    ih.simpan_draf(A, "ruang", {"daftar": [{"jenis": "Ruang %d" % i} for i in range(30)]}, "adma")
    versi = ih.baca(A, "ruang")["versi"]
    r = unggah(ka, "ruang", "daftar", xlsx([["Jenis ruang"]] + [["Lain %d" % i] for i in range(21)]))
    teks = r.get_data(as_text=True)
    cek('value="tambah" disabled' in teks and 'value="ganti" checked' in teks and "melebihi batas 50 baris" in teks,
        "pratinjau: tambah dimatikan bila melebihi batas")
    r = terapkan(ka, "ruang", token_dari(r), "tambah")
    cek("Impor dibatalkan" in ka.get(r.headers["Location"]).get_data(as_text=True) and ih.baca(A, "ruang")["versi"] == versi,
        "tambah melebihi batas ditolak, draf tetap")
    for nama_berkas, isi, harus in (("lama.xls", b"\xd0\xcf\x11\xe0" + b"0" * 50, "Format .xls lama"),
                                    ("isi.xlsx", xlsx([["Semester"], [1]]), "kolom wajib tidak ada: Mata kuliah, SKS")):
        r = unggah(ka, "kurikulum", "mata_kuliah", isi, nama_berkas)
        cek(r.status_code == 400 and harus in r.get_data(as_text=True), "unggahan ditolak dengan alasan %r" % harus)
    r = kirim(ka, URL + "/impor", tabel="mata_kuliah")
    cek(r.status_code == 400 and "pilih berkas Excel" in r.get_data(as_text=True), "tanpa berkas")
    cek(kirim(ka, URL + "/impor", tabel="dokumen", berkas=(io.BytesIO(b"x"), "x.csv")).status_code == 400, "tabel tidak dikenal")
except Exception as e:                 # a step that went wrong stops the suite; say where instead of only a traceback
    import traceback
    traceback.print_exc()
    gagal.insert(0, "uji berhenti di tengah (%s: %s); pemeriksaan sebelumnya:" % (type(e).__name__, e))
finally:
    shutil.rmtree(KERJA, ignore_errors=True)

for g_ in gagal:
    print("GAGAL", g_)
print("uji_impor_tabel: %s" % ("LULUS" if not gagal else "%d gagal" % len(gagal)))
sys.exit(1 if gagal else 0)
