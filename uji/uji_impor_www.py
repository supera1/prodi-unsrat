# -*- coding: utf-8 -*-
"""Step 7 of phase 1: www's data uploaded by the superadmin (K-P18: platform/impor_www.py and the "Data www" page),
and the starting copy of a new VM (platform/data-awal/www/).

- the Excel export of www's superadmin, written here in www's layout (and, when the www checkout sits next to this
  repository, made by www's own export code from www's real data): read back without a single change; changed cells,
  a programme new in www, a programme left out, a missing column, a missing CPL sheet, every kind of bad row;
- the file checks: not .xlsx, macros, zip bombs, sizes, too many files, two sources of the same data;
- www's JSON files, recognised by their content; what is not uploaded is taken from the current copy;
- the admin pages: upload, preview, apply once, only by the uploader, within 30 minutes, never onto a copy that
  changed after the upload; gambling text refused; a refused upload leaves nothing behind; the dashboard;
- an empty platform: the starting copy and the registry from the slug list (--isi-awal), or a first upload.

Who may open these pages at all is tested in uji_admin.py.
Needs the packages of platform/admin/requirements.txt: run with .venv/bin/python.
"""
from __future__ import print_function

import datetime
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

AKAR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for sub in ("platform", "platform/admin", "platform/bangun", "platform/pemantau"):
    sys.path.insert(0, os.path.join(AKAR, *sub.split("/")))

try:
    import flask  # noqa: F401
    import openpyxl
except ImportError:
    print("uji_impor_www: paket admin belum terpasang; jalankan dengan .venv/bin/python (lihat README)")
    sys.exit(1)

KERJA = tempfile.mkdtemp(prefix="prodi-uji-impor-")
for nama, sub in (("PRODI_DATA", "data"), ("PRODI_RIWAYAT", "riwayat"), ("PRODI_RAHASIA", "rahasia"), ("PRODI_LOG", "log"),
                  ("WWW_DATA", "www"), ("PRODI_PUBLIK", "publik"), ("PRODI_NGINX", "nginx"), ("PRODI_UNGGAHAN", "unggahan")):
    os.environ[nama] = os.path.join(KERJA, sub)
os.environ["PRODI_MANIFES"] = os.path.join(KERJA, "manifes.json")
for nama in ("PRODI_MODE", "PRODI_PENYIMPANAN"):
    os.environ.pop(nama, None)

import akun  # noqa: E402
import app as app_mod  # noqa: E402
import audit  # noqa: E402
import impor_www  # noqa: E402
import lokasi  # noqa: E402
import registri  # noqa: E402
import tarik_www  # noqa: E402

WWW_SEBELAH = os.path.join(os.path.dirname(AKAR), "website-unsrat-2026", "site")
IP = "103.84.116.10"
A, B, C, D = "55201", "20201", "21201", "44201"
SANDI = {"superuji": "Biru laut 7 menara pagi", "superdua": "Hujan 9 jendela kayu tua"}
FMIPA = "Fakultas Matematika dan Ilmu Pengetahuan Alam"
gagal = []


def cek(kondisi, pesan):
    if not kondisi:
        gagal.append(pesan)


def harus_gagal(fungsi, pesan, *args, **kw):
    """The GagalImpor of the call, or a recorded failure when the call goes through."""
    try:
        fungsi(*args, **kw)
    except impor_www.GagalImpor as e:
        return e
    gagal.append(pesan)
    return impor_www.GagalImpor([])


def ada_pesan(e, potongan):
    return any(potongan in m for m in e.masalah)


def prodi(kode, nama, **lain):
    p = {"kode_pddikti": kode, "nama": nama, "nama_en": "", "jenjang": "S1", "gelar": "S.T.", "sks": 144, "unit_kode": "ft",
         "unit": "Fakultas Teknik", "website": "", "surel": "x@unsrat.ac.id", "bidang_kerja": "Bidang.", "bidang_kerja_en": "",
         "akreditasi": "Baik", "lembaga_akreditasi": "LAM-Teknik", "berlaku_hingga": "2028-01-31", "internasional": "",
         "slug": "s1-" + nama.lower().replace(" ", "-"), "slug_en": "",
         "cpl": [{"kode": "CPL 1", "isi": "Capaian %s." % nama, "isi_en": "", "kurikulum": "2025/2026"}]}
    p.update(lain)
    return p


WWW = {"prodi": {"tampilkan_berlaku_hingga": True, "program": [
           prodi(A, "Teknik Informatika", akreditasi="Unggul"), prodi(B, "Teknik Elektro", akreditasi="belum"),
           prodi(C, "Teknik Mesin")]},
       "unit": {"unit": [{"kode": "ft", "nama": "Fakultas Teknik", "bendera": {"dasar": "#000099", "tulisan": "#FFFF00"}},
                         {"kode": "fmipa", "nama": FMIPA, "bendera": {"dasar": "#FFFFFF", "tulisan": "#000000"}}]},
       "kalender": {"semester": []},
       "situs": {"identitas": {"nama": {"id": "Universitas Sam Ratulangi"}}, "kontak": {}, "media_sosial": [],
                 "utilitas": {"id": []}}}


def tulis_www(isi):
    os.makedirs(os.environ["WWW_DATA"], exist_ok=True)
    for nama, data in isi.items():
        with io.open(os.path.join(os.environ["WWW_DATA"], nama + ".json"), "w", encoding="utf-8") as f:
            json.dump(data, f)


LABEL_PRODI = [l for _k, l in impor_www.KOLOM_PRODI] + ["Terjemahan Inggris diperiksa"]


def excel(program, cpl=True, buang=(), atas=0, cpl_tambahan=()):
    """An .xlsx in the layout of www's export (site/admin/impor.py, Prodi.ekspor): sheets "Program studi", "CPL",
    "Petunjuk", "Pilihan"; dates as date cells, "belum" as an empty cell. buang: column labels left out; atas: title
    rows above the header row; cpl_tambahan: extra rows of the CPL sheet."""
    kunci = {l: k for k, l in impor_www.KOLOM_PRODI}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Program studi"
    for i in range(atas):
        ws.append(["Templat program studi UNSRAT"] if i == 0 else [])
    label = [l for l in LABEL_PRODI if l not in buang]
    ws.append(label)
    for p in program:
        baris = []
        for l in label:
            v = p.get(kunci[l]) if l in kunci else None
            if kunci.get(l) == "akreditasi" and v == "belum":
                v = None
            if kunci.get(l) == "berlaku_hingga" and re.match(r"^\d{4}-\d{2}-\d{2}$", str(v or "")):
                v = datetime.datetime.strptime(v, "%Y-%m-%d")
            baris.append(None if v == "" else v)
        ws.append(baris)
    if cpl:
        wc = wb.create_sheet("CPL")
        wc.append([l for _k, l in impor_www.KOLOM_CPL])
        for p in program:
            for c in p.get("cpl") or []:
                wc.append([p["kode_pddikti"], c["kode"], c["isi"], c.get("isi_en") or None, c.get("kurikulum") or None])
        for b in cpl_tambahan:
            wc.append(list(b))
    wb.create_sheet("Petunjuk").append(["Lembar 'Program studi': satu baris per program studi."])
    wb.create_sheet("Pilihan").append(["Jenjang", "Kode unit", "Nama unit", "Akreditasi", "Lembaga (saran)"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def json_bytes(isi, bom=False):
    return (u"﻿" if bom else u"").encode("utf-8") + json.dumps(isi).encode("utf-8")


nomor = [0]


def siapkan(berkas):
    nomor[0] += 1
    folder = os.path.join(KERJA, "kiriman-uji", "k%d" % nomor[0])
    return folder, impor_www.siapkan_kiriman(berkas, folder, pemilik=1)


def platform_baru(nama):
    """Point the platform at an empty data folder."""
    for env, sub in (("PRODI_DATA", "data"), ("PRODI_RIWAYAT", "riwayat")):
        os.environ[env] = os.path.join(KERJA, nama, sub)


try:
    tulis_www(WWW)
    tarik_www.terapkan("penguji")
    program = [dict(p) for p in WWW["prodi"]["program"]]

    # --- the unchanged export reads back without a single change, also with title rows above the header and with
    # codes and SKS typed as numbers
    for label, isi in (("ekspor apa adanya", excel(program)), ("judul di atas kepala tabel", excel(program, atas=3)),
                       ("kode dan SKS angka", excel([dict(p, kode_pddikti=int(p["kode_pddikti"]), sks=float(p["sks"]))
                                                     for p in program]))):
        folder, info = siapkan([("unsrat-prodi-2026-10-26.xlsx", isi)])
        laporan = tarik_www.pratinjau(folder)["laporan"]
        cek(not tarik_www.ada_perubahan(laporan), "%s: tidak ada perubahan: %r" % (label, {k: v for k, v in laporan.items() if v}))
        cek(info["dari_salinan"] == ["unit", "kalender", "situs"] and not info["peringatan"], "%s: %r" % (label, info))
    cek(info["label"] == "unggahan superadmin: unsrat-prodi-2026-10-26.xlsx", "label sumber: %r" % info["label"])
    cek(sorted(os.listdir(folder)) == ["kalender.json", "kiriman.json", "prodi.json", "situs.json", "unit.json"],
        "isi folder unggahan: %r" % os.listdir(folder))
    cek(oct(os.stat(folder).st_mode & 0o777) == "0o700", "folder unggahan hanya untuk pemiliknya")

    # --- changed cells, a programme new in www (unit code in capitals), a programme left out
    ubah = [dict(program[0], akreditasi="baik sekali", sks=150, berlaku_hingga="2030-06-30", nama_en="Informatics"),
            dict(program[1], cpl=[{"kode": "CPL 1", "isi": "Capaian baru.", "isi_en": "New outcome.", "kurikulum": "2025"},
                                  {"kode": "CPL 2", "isi": "Capaian kedua.", "isi_en": "", "kurikulum": ""}]),
            prodi(D, "Matematika", unit_kode="FMIPA", unit="", akreditasi="belum", sks=None, gelar="S.Mat.", slug="",
                  berlaku_hingga="", cpl=[])]
    folder, info = siapkan([("prodi.xlsx", excel(ubah))])
    laporan = tarik_www.pratinjau(folder)["laporan"]
    cek(laporan["baru"] == [D] and laporan["hilang"] == [C] and sorted(laporan["berubah"]) == [B, A],
        "pratinjau perubahan: %r" % laporan)
    cek("akreditasi: 'Unggul' -> 'Baik Sekali'" in laporan["berubah"][A] and "sks: 144 -> 150" in laporan["berubah"][A]
        and "berlaku_hingga: '2028-01-31' -> '2030-06-30'" in laporan["berubah"][A], "rincian A: %r" % laporan["berubah"][A])
    cek("CPL diubah: CPL 1" in laporan["berubah"][B] and "CPL ditambah: CPL 2" in laporan["berubah"][B], "rincian B")
    cek(laporan["hilang_terlalu_banyak"], "1 dari 3 prodi hilang melewati batas")
    try:
        tarik_www.terapkan("penguji", folder=folder, keterangan=info["label"])
        gagal.append("terlalu banyak prodi hilang harus ditolak tanpa paksa")
    except tarik_www.GagalTarik:
        pass
    tarik_www.terapkan("penguji", folder=folder, keterangan=info["label"], paksa=True)
    w = registri.warisan_per_kode()
    cek(w[A]["akreditasi"] == "Baik Sekali" and w[A]["sks"] == 150 and w[A]["nama_en"] == "Informatics", "A: %r" % w[A])
    cek(w[A]["slug"] == "s1-teknik-informatika" and w[A]["gelar"] == "S.T.", "kolom yang tidak berubah tetap")
    cek(w[D]["unit_kode"] == "fmipa" and w[D]["unit"] == FMIPA and w[D]["akreditasi"] == "belum" and w[D]["sks"] is None
        and w[D]["slug"] == "" and w[D]["cpl"] == [], "prodi baru dari Excel: %r" % w[D])
    cek(w[C].get("hilang_dari_www"), "prodi yang tidak ada di berkas ditandai hilang")
    cek(lokasi.baca("warisan/prodi")["tampilkan_berlaku_hingga"] is True, "pengaturan www yang tidak ada di Excel tetap")
    cek(lokasi.baca("warisan/meta")["sumber"] == "unggahan superadmin: prodi.xlsx", "meta mencatat sumber")
    sekarang = [dict(w[A]), dict(w[B]), dict(w[D])]

    # --- a column left out keeps its values; a missing CPL sheet keeps every CPL
    folder, info = siapkan([("prodi.xlsx", excel([dict(p, gelar="DIUBAH", nama_en="Informatika EN" if p is sekarang[0] else p["nama_en"])
                                                  for p in sekarang], buang=("Gelar",), cpl=False))])
    cek("Kolom 'Gelar' tidak ada di berkas; nilainya tidak diubah." in info["peringatan"]
        and "Lembar CPL tidak ada di berkas; CPL tidak diubah." in info["peringatan"], "peringatan: %r" % info["peringatan"])
    laporan = tarik_www.pratinjau(folder)["laporan"]
    cek(list(laporan["berubah"]) == [A] and laporan["berubah"][A] == ["nama_en: 'Informatics' -> 'Informatika EN'"],
        "hanya kolom yang ada yang berubah: %r" % laporan["berubah"])

    # --- an accreditation value www may add later is kept as it is, with a warning
    folder, info = siapkan([("prodi.xlsx", excel([dict(sekarang[0], akreditasi="Unggul Internasional")] + sekarang[1:]))])
    cek(any("akreditasi 'Unggul Internasional' belum dikenal" in m for m in info["peringatan"]), "akreditasi baru diperingatkan")
    cek(json.load(open(os.path.join(folder, "prodi.json")))["program"][0]["akreditasi"] == "Unggul Internasional",
        "akreditasi baru disimpan apa adanya")
    for teks, nilai in (("", "belum"), ("Belum dilengkapi", "belum"), ("terakreditasi sementara", "Terakreditasi Sementara"),
                        ("TERAKREDITASI", "TERAKREDITASI"), ("Menunggu akreditasi prodi baru", "Menunggu Akr. PS Baru")):
        folder, info = siapkan([("prodi.xlsx", excel([dict(sekarang[0], akreditasi=teks or "belum")] + sekarang[1:]))])
        hasil = json.load(open(os.path.join(folder, "prodi.json")))["program"][0]["akreditasi"]
        cek(hasil == nilai and not info["peringatan"], "akreditasi %r -> %r, harus %r" % (teks, hasil, nilai))

    # --- every kind of bad row, all reported at once with their row
    buruk = sekarang + [prodi("55-201", "Kode Salah"), dict(sekarang[0]), prodi("90201", "Jenjang Salah", jenjang="D3"),
                        prodi("90202", "Unit Salah", unit_kode="fkik"), prodi("90203", "SKS Salah", sks="seratus"),
                        prodi("90204", "Tanggal Salah", berlaku_hingga="31/31/2030"), prodi("90205", "")]
    e = harus_gagal(impor_www.siapkan_kiriman, "baris buruk harus ditolak",
                    [("prodi.xlsx", excel(buruk, cpl_tambahan=[("99999", "CPL 1", "Isi.", None, None), (A, "CPL 9", None, None, None)]))],
                    os.path.join(KERJA, "kiriman-uji", "buruk"))
    for potongan in ("baris 5: kode PDDIKTI '55-201' tidak sah", "baris 6: kode 55201 muncul lebih dari sekali",
                     "baris 7, kolom Jenjang: jenjang 'D3' tidak dikenal", "baris 8, kolom Kode unit: kode unit 'fkik'",
                     "baris 9, kolom SKS: SKS 'seratus' harus angka", "baris 10, kolom Berlaku hingga: tanggal '31/31/2030'",
                     "baris 11, kolom Nama program studi: wajib diisi", "tidak ada di lembar Program studi",
                     "Kode CPL dan Isi CPL wajib diisi"):
        cek(ada_pesan(e, potongan), "masalah %r harus dilaporkan: %r" % (potongan, e.masalah))
    cek(not os.path.exists(os.path.join(KERJA, "kiriman-uji", "buruk")), "unggahan yang ditolak tidak meninggalkan folder")
    wb = openpyxl.Workbook()
    wb.active.append(["Kode PDDIKTI", "Nama program studi", "Jenjang"])
    wb.active.append([A, "Teknik Informatika", "S1"])
    buf = io.BytesIO()
    wb.save(buf)
    e = harus_gagal(siapkan, "tanpa kolom Kode unit harus ditolak", [("prodi.xlsx", buf.getvalue())])
    cek(ada_pesan(e, "judul kolom 'Kode PDDIKTI', 'Nama program studi', 'Jenjang', 'Kode unit' tidak ditemukan"), "%r" % e.masalah)

    # --- gambling text passes the reading but never the checks of tarik_www
    folder, info = siapkan([("prodi.xlsx", excel([dict(sekarang[0], bidang_kerja="Daftar slot gacor di sini")] + sekarang[1:]))])
    try:
        tarik_www.pratinjau(folder)
        gagal.append("teks judi dari Excel harus ditolak")
    except tarik_www.GagalTarik as e:
        cek(any("teks judi" in m for m in e.masalah), "teks judi dilaporkan: %r" % e.masalah)

    # --- the file checks
    baik = excel(sekarang)

    def zip_biasa(isi):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for nama, data in isi:
                z.writestr(nama, data)
        return buf.getvalue()

    def zip_dengan(tambahan):
        """The good workbook with extra entries."""
        sumber = zipfile.ZipFile(io.BytesIO(baik))
        return zip_biasa([(i.filename, sumber.read(i.filename)) for i in sumber.infolist()] + list(tambahan))
    for label, berkas, potongan in (
            ("tanpa berkas", [], "Pilih berkas"),
            ("xls lama", [("prodi.xls", baik)], "format .xls tidak diterima"),
            ("xlsm", [("prodi.xlsm", baik)], "format .xlsm tidak diterima"),
            ("csv", [("prodi.csv", b"a,b")], "format .csv tidak diterima"),
            ("berkas lain", [("prodi.pdf", b"%PDF")], "hanya berkas Excel (.xlsx) atau JSON"),
            ("bukan zip", [("prodi.xlsx", b"bukan excel sama sekali")], "bukan file Excel .xlsx yang sah"),
            ("zip rusak", [("prodi.xlsx", b"PK\x03\x04" + b"\x00" * 100)], "rusak atau bukan .xlsx"),
            ("zip bukan Excel", [("prodi.xlsx", zip_biasa([("catatan.txt", "bukan lembar kerja")]))],
             "bukan file Excel .xlsx yang sah"),
            ("makro", [("prodi.xlsx", zip_dengan([("xl/vbaProject.bin", b"\x00" * 64)]))], "makro"),
            ("bom zip", [("prodi.xlsx", zip_dengan([("xl/media/besar.bin", b"\x00" * (20 * 1024 * 1024))]))],
             "terlalu besar atau tidak wajar"),
            ("lebih dari 5 MB", [("prodi.xlsx", b"PK\x03\x04" + b"\x00" * impor_www.MAKS_UKURAN)], "maksimal 5 MB"),
            ("JSON lebih dari 5 MB", [("prodi.json", b" " * (impor_www.MAKS_UKURAN + 1))], "maksimal 5 MB"),
            ("terlalu banyak berkas", [("u%d.json" % i, b"{}") for i in range(impor_www.MAKS_BERKAS + 1)], "Paling banyak"),
            ("dua Excel", [("a.xlsx", baik), ("b.xlsx", baik)], "Unggah satu berkas Excel saja"),
            ("Excel dan prodi.json", [("a.xlsx", baik), ("prodi.json", json_bytes(WWW["prodi"]))], "diunggah dua kali"),
            ("dua unit.json", [("unit.json", json_bytes(WWW["unit"])), ("unit (1).json", json_bytes(WWW["unit"]))],
             "sama-sama berisi data fakultas"),
            ("JSON tak dikenal", [("lain.json", b'{"a": 1}')], "bukan berkas data www yang dikenal"),
            ("JSON ganda makna", [("lain.json", b'{"program": [], "unit": []}')], "bukan berkas data www yang dikenal"),
            ("JSON rusak", [("unit.json", b"{bukan json")], "bukan berkas JSON yang sah"),
            ("JSON bersarang dalam", [("unit.json", b"[" * 100000 + b"]" * 100000)], "bukan berkas JSON yang sah"),
            ("JSON bukan UTF-8", [("unit.json", b'{"unit": ["\xff"]}')], "bukan berkas JSON yang sah")):
        e = harus_gagal(siapkan, "%s harus ditolak" % label, berkas)
        cek(ada_pesan(e, potongan), "%s: pesan %r tidak ada di %r" % (label, potongan, e.masalah))

    # --- JSON files: recognised by content whatever their name, with a byte-order mark too; an Excel file may use a
    # faculty that only the uploaded unit.json has
    unit_baru = {"unit": WWW["unit"]["unit"] + [{"kode": "fpik", "nama": "Fakultas Perikanan dan Ilmu Kelautan"}]}
    folder, info = siapkan([("unduhan (3).json", json_bytes(unit_baru, bom=True)),
                            ("prodi.xlsx", excel(sekarang + [prodi("54201", "Ilmu Kelautan", unit_kode="fpik")]))])
    cek([b["jenis"] for b in info["berkas"]] == ["prodi", "unit"] and info["dari_salinan"] == ["kalender", "situs"], "%r" % info)
    laporan = tarik_www.pratinjau(folder)["laporan"]
    cek(laporan["unit_berubah"] and laporan["baru"] == ["54201"] and not laporan["kalender_berubah"], "unit.json + Excel: %r" % laporan)
    cek(json.load(open(os.path.join(folder, "prodi.json")))["program"][-1]["unit"] == "Fakultas Perikanan dan Ilmu Kelautan",
        "nama unit prodi baru dari unit.json yang diunggah")
    folder, info = siapkan([("kalender.json", json_bytes({"semester": [{"nama": "Ganjil"}]})),
                            ("situs.json", json_bytes(dict(WWW["situs"], kontak={"telepon": "(0000) 000000"})))])
    laporan = tarik_www.pratinjau(folder)["laporan"]
    cek(laporan["kalender_berubah"] and laporan["universitas_berubah"] and not laporan["berubah"] and not laporan["unit_berubah"],
        "hanya kalender dan identitas yang berubah: %r" % laporan)

    # --- the fingerprint of the copy an upload was made from
    folder, info = siapkan([("prodi.xlsx", baik)])
    cek(info["dasar"] == impor_www.sidik_dasar(), "sidik dasar dicatat")
    f2, i2 = siapkan([("unit.json", json_bytes(unit_baru))])
    tarik_www.terapkan("penguji", folder=f2)
    cek(impor_www.sidik_dasar() != info["dasar"], "sidik dasar berubah setelah salinan berubah")

    # --- the real www export, made by www's own code, reads back without a change (developer's machine only)
    if os.path.isfile(os.path.join(WWW_SEBELAH, "admin", "impor.py")):
        path_xlsx = os.path.join(KERJA, "ekspor-www.xlsx")
        skrip = ("import json, sys; sys.path.insert(0, sys.argv[1]); import impor; "
                 "d = json.load(open(sys.argv[2], encoding='utf-8')); "
                 "open(sys.argv[3], 'wb').write(impor.ke_bytes(impor.Prodi().ekspor(d)))")
        hasil = subprocess.run([sys.executable, "-I", "-c", skrip, os.path.join(WWW_SEBELAH, "admin"),
                                os.path.join(WWW_SEBELAH, "mockup", "data", "prodi.json"), path_xlsx],
                               capture_output=True, text=True, timeout=120)
        cek(hasil.returncode == 0, "ekspor www: %s" % hasil.stderr[-500:])
        platform_baru("asli")
        tarik_www.terapkan("penguji", os.path.join(WWW_SEBELAH, "mockup", "data"))
        with open(path_xlsx, "rb") as f:
            folder, info = siapkan([("unsrat-prodi.xlsx", f.read())])
        laporan = tarik_www.pratinjau(folder)["laporan"]
        cek(not tarik_www.ada_perubahan(laporan) and not info["peringatan"],
            "ekspor www asli harus terbaca tanpa perubahan: %r" % {k: v for k, v in laporan.items() if v and k != "kekurangan"})
        print("ekspor Excel asli dari www: %d prodi terbaca tanpa perubahan" % len(lokasi.baca("warisan/prodi")["program"]))
        # the starting copy of new VMs against www's data now
        awal = {p["kode_pddikti"]: p for p in impor_www.salinan_sekarang("prodi")["program"]}
        with io.open(os.path.join(impor_www.DATA_AWAL, "prodi.json"), encoding="utf-8") as f:
            if {p["kode_pddikti"]: p for p in json.load(f)["program"]} != awal:
                print("PERHATIAN platform/data-awal/www/ sudah berbeda dari data www; perbarui dengan "
                      "impor_www.py --tulis-data-awal (lihat docs/deploy.md)")
    else:
        print("checkout www tidak ada di sebelah repo: uji ekspor Excel asli dilewati")

    # --- an empty platform: the starting copy and the registry from the approved slug list
    platform_baru("kosong")
    e = harus_gagal(siapkan, "Excel saja di platform kosong harus ditolak", [("prodi.xlsx", baik)])
    cek(ada_pesan(e, "Data fakultas belum ada di platform"), "%r" % e.masalah)
    e = harus_gagal(siapkan, "tanpa kalender dan situs di platform kosong harus ditolak",
                    [("prodi.xlsx", baik), ("unit.json", json_bytes(WWW["unit"]))])
    cek(ada_pesan(e, "Platform belum punya kalender akademik") and ada_pesan(e, "Platform belum punya identitas"), "%r" % e.masalah)
    asal = impor_www.asal_data_awal()
    cek(asal.get("jumlah_prodi") == 97 and re.match(r"^\d{4}-\d{2}-\d{2}$", asal.get("diambil", "")), "asal data awal: %r" % asal)
    pesan = impor_www.isi_awal()
    cek(len(pesan) == 2 and "97 prodi" in pesan[0] and "97 prodi" in pesan[1], "isi awal: %r" % pesan)
    cek(len(registri.baca()["prodi"]) == 97 and len(os.listdir(os.path.join(lokasi.data(), "prodi"))) == 97,
        "registri dan pengaturan 97 prodi dibuat")
    cek(lokasi.baca("warisan/meta")["sumber"] == "data awal repo (salinan www %s)" % asal["diambil"], "sumber data awal")
    cek(impor_www.isi_awal() == [], "isi awal kedua kali tidak melakukan apa pun")
    cek(impor_www.main(["--isi-awal"]) == 0, "baris perintah --isi-awal")
    # the starting copy, applied and written again, is the same (never written into the repository here)
    tulis_ke = os.path.join(KERJA, "data-awal-uji")
    cek(impor_www.tulis_data_awal(tulis_ke, dari="uji") == 97, "tulis data awal")
    for nama in impor_www.JENIS:
        with open(os.path.join(tulis_ke, nama + ".json"), "rb") as f1, open(os.path.join(impor_www.DATA_AWAL, nama + ".json"), "rb") as f2:
            cek(json.load(f1) == json.load(f2), "data awal %s.json bolak-balik tanpa perubahan" % nama)
    platform_baru("tanpa-salinan")
    harus_gagal(impor_www.tulis_data_awal, "tulis data awal tanpa salinan harus gagal", os.path.join(KERJA, "data-awal-kosong"))
    cek(not os.path.exists(os.path.join(KERJA, "data-awal-kosong")), "tanpa salinan tidak ada yang ditulis")

    # --------------------------------------------------------------- the admin pages
    platform_baru("admin")
    tarik_www.terapkan("penguji")                      # WWW_DATA: A, B, C
    path_csv = os.path.join(KERJA, "slug.csv")
    with io.open(path_csv, "w", encoding="utf-8") as f:
        f.write(",".join(registri.slug_mod.KOLOM) + "\n%s,S1,Teknik Informatika,FT,s1informatika,x,x\n"
                "%s,S1,Teknik Elektro,FT,s1elektro,x,x\n" % (A, B))
    registri.siapkan_dari_csv("penguji", path_csv)
    shutil.rmtree(os.environ["WWW_DATA"])               # the server has no WWW_DATA folder: uploads only
    ids = {}
    for nama in SANDI:
        ids[nama], sementara = akun.buat(None, nama, nama.capitalize() + " Uji", "super")
        akun.ganti_sandi(ids[nama], sementara, SANDI[nama])
    app = app_mod.buat_app({"SECRET_KEY": "k" * 64, "ADMIN_IP_BOLEH": "103.84.116.0/24", "PRODI_BANGUN": "sinkron"})

    def klien(nama):
        c = app.test_client()
        c.environ_base["REMOTE_ADDR"] = IP
        c.get("/masuk")
        with c.session_transaction() as s:
            token = s["csrf"]
        r = c.post("/masuk", data={"csrf": token, "pengguna": nama, "sandi": SANDI[nama]})
        cek(r.status_code == 302, "%s masuk" % nama)
        return c

    def kirim(c, url, berkas=None, **data):
        with c.session_transaction() as s:
            data["csrf"] = s.get("csrf", "")
        if berkas is not None:
            data["berkas"] = [(io.BytesIO(isi), nama) for nama, isi in berkas]
            return c.post(url, data=data, content_type="multipart/form-data")
        return c.post(url, data=data)

    def unggah(c, berkas):
        r = kirim(c, "/super/tarik/unggah", berkas)
        m = re.search(r"kiriman=([A-Za-z0-9_-]+)", r.headers.get("Location", ""))
        return r, m.group(1) if m else None

    def isi_formulir(teks):
        m = re.search(r'name="sidik" value="([0-9a-f]*)"', teks)
        return m.group(1) if m else None

    def kiriman_ada():
        akar = os.path.join(lokasi.rahasia(), "kiriman")
        return sorted(os.listdir(akar)) if os.path.isdir(akar) else []

    su, sd = klien("superuji"), klien("superdua")
    teks = su.get("/super/tarik").get_data(as_text=True)
    cek("Unggah data dari www" in teks and 'enctype="multipart/form-data"' in teks and "Pratinjau" not in teks,
        "tanpa WWW_DATA halaman hanya berisi formulir unggah")
    awal = registri.warisan_per_kode()
    baris = [dict(awal[k], akreditasi="Baik Sekali" if k == A else awal[k]["akreditasi"]) for k in (A, B, C)]
    r, token = unggah(su, [("unsrat-prodi-2026-10-26.xlsx", excel(baris))])
    cek(r.status_code == 302 and token, "unggahan diterima: %s" % r.status_code)
    teks = su.get("/super/tarik?kiriman=%s" % token).get_data(as_text=True)
    sidik = isi_formulir(teks)
    cek("Pratinjau unggahan" in teks and "unsrat-prodi-2026-10-26.xlsx" in teks and "Baik Sekali" in teks
        and 'name="kiriman" value="%s"' % token in teks and sidik, "pratinjau unggahan")
    cek("tetap memakai salinan sekarang: fakultas, kalender akademik, identitas universitas" in teks, "yang tidak diunggah")
    # only the uploader
    r = sd.get("/super/tarik?kiriman=%s" % token)
    cek(r.status_code == 302 and "kiriman" not in r.headers["Location"], "superadmin lain tidak bisa membuka unggahan")
    cek(kirim(sd, "/super/tarik/terapkan", sidik=sidik, kiriman=token).status_code == 302
        and registri.warisan_per_kode()[A]["akreditasi"] == "Unggul", "superadmin lain tidak bisa menerapkan")
    # an outdated preview
    r = kirim(su, "/super/tarik/terapkan", sidik="0" * 32, kiriman=token)
    cek(r.status_code == 302 and "kiriman=%s" % token in r.headers["Location"]
        and registri.warisan_per_kode()[A]["akreditasi"] == "Unggul", "pratinjau usang tidak diterapkan")
    # applied once
    r = kirim(su, "/super/tarik/terapkan", sidik=sidik, kiriman=token)
    cek(r.status_code == 302 and registri.warisan_per_kode()[A]["akreditasi"] == "Baik Sekali", "unggahan diterapkan")
    cek(app.antrean.status["ok"] and os.path.isfile(os.path.join(lokasi.publik(), "s1informatika", "index.html")),
        "situs dibangun ulang: %s" % app.antrean.status["pesan"])
    cek(token not in kiriman_ada(), "folder unggahan dihapus setelah diterapkan")
    cek(any(x["aksi"] == "tarik-data-www" and x.get("sumber") == "unggahan superadmin: unsrat-prodi-2026-10-26.xlsx"
            and x.get("oleh") == "superuji" for x in audit.baca()), "penerapan tercatat dengan sumbernya")
    cek(any(x["aksi"] == "data-www-diunggah" and x.get("berkas") == ["unsrat-prodi-2026-10-26.xlsx"] for x in audit.baca()),
        "unggahan tercatat")
    r = kirim(su, "/super/tarik/terapkan", sidik=sidik, kiriman=token)
    cek(r.status_code == 302 and su.get("/super/tarik?kiriman=%s" % token).status_code == 302, "unggahan hanya sekali pakai")
    cek("tidak berlaku" in su.get("/super/tarik").get_data(as_text=True), "pesan unggahan tidak berlaku")
    teks = su.get("/").get_data(as_text=True)
    cek("sumber: unggahan superadmin: unsrat-prodi-2026-10-26.xlsx" in teks and "(0 hari lalu)" in teks, "dasbor: sumber dan umur")

    # an upload made before another one was applied is not applied over it
    baris = [dict(registri.warisan_per_kode()[k]) for k in (A, B, C)]
    _r, lama = unggah(su, [("prodi.xlsx", excel([dict(baris[0], nama_en="Lama")] + baris[1:]))])
    _r, baru = unggah(su, [("unit.json", json_bytes(unit_baru))])
    teks = su.get("/super/tarik?kiriman=%s" % baru).get_data(as_text=True)
    cek(kirim(su, "/super/tarik/terapkan", sidik=isi_formulir(teks), kiriman=baru).status_code == 302
        and lokasi.baca("warisan/unit") == unit_baru, "unggahan unit.json diterapkan")
    teks = su.get("/super/tarik?kiriman=%s" % lama).get_data(as_text=True)
    cek("berubah sejak berkas ini diunggah" in teks and 'name="kiriman"' not in teks, "unggahan lama ditandai usang")
    r = kirim(su, "/super/tarik/terapkan", sidik=isi_formulir(teks) or "x", kiriman=lama)
    cek(r.status_code == 302 and registri.warisan_per_kode()[A]["nama_en"] != "Lama" and lama not in kiriman_ada()
        and lokasi.baca("warisan/unit") == unit_baru, "unggahan usang tidak menimpa perubahan sesudahnya")

    # a refused upload: the problems, nothing left behind
    sebelum = kiriman_ada()
    r = kirim(su, "/super/tarik/unggah", [("prodi.xlsx", excel([prodi(A, "Teknik Informatika", jenjang="D4")]))])
    teks = r.get_data(as_text=True)
    cek(r.status_code == 400 and "Berkas unggahan tidak dapat dipakai" in teks and "jenjang &#39;D4&#39; tidak dikenal" in teks,
        "unggahan buruk ditolak dengan masalahnya")
    cek(kiriman_ada() == sebelum, "unggahan yang ditolak tidak meninggalkan folder")
    cek(any(x["aksi"] == "data-www-unggahan-ditolak" for x in audit.baca()), "penolakan tercatat")
    r = kirim(su, "/super/tarik/unggah", [])
    cek(r.status_code == 400 and "Pilih berkas" in r.get_data(as_text=True), "unggah tanpa berkas")

    # gambling text in an uploaded file: shown as a problem, never applied
    r, token = unggah(su, [("prodi.xlsx", excel([dict(baris[0], bidang_kerja="situs s1ot g4cor")] + baris[1:]))])
    teks = su.get("/super/tarik?kiriman=%s" % token).get_data(as_text=True)
    cek("Data www tidak dapat ditarik" in teks and "teks judi" in teks and 'name="sidik"' not in teks, "judi di unggahan")
    r = kirim(su, "/super/tarik/terapkan", sidik="x", kiriman=token)
    cek(r.status_code == 400 and "situs s1ot g4cor" not in json.dumps(lokasi.baca("warisan/prodi")), "judi tidak diterapkan")

    # 30 minutes; strange tokens; old uploads are cleared away
    r, token = unggah(su, [("prodi.xlsx", excel(baris))])
    path = os.path.join(lokasi.rahasia(), "kiriman", token, "kiriman.json")
    info = json.load(open(path))
    lokasi.tulis_json(path, dict(info, waktu=time.time() - 31 * 60))
    cek(su.get("/super/tarik?kiriman=%s" % token).status_code == 302, "unggahan lebih dari 30 menit tidak berlaku")
    for aneh in ("../../rahasia", "x" * 10, token + "/../" + token):
        cek(su.get("/super/tarik", query_string={"kiriman": aneh}).status_code == 302, "token aneh %r ditolak" % aneh)
    os.utime(os.path.join(lokasi.rahasia(), "kiriman", token), (time.time() - 31 * 60,) * 2)
    unggah(su, [("prodi.xlsx", excel(baris))])
    cek(token not in kiriman_ada(), "unggahan kedaluwarsa dibersihkan")

    # the dashboard reminds the superadmin of old data
    meta = lokasi.baca("warisan/meta")
    lokasi.tulis_json(lokasi.path_data("warisan/meta"), dict(meta, ditarik=(datetime.datetime.now() - datetime.timedelta(days=40)).isoformat()))
    cek("lebih dari 30 hari" in su.get("/").get_data(as_text=True), "dasbor mengingatkan data www yang lama")

    # a new VM whose starting copy was not applied: the first upload brings the data and the registry
    platform_baru("pertama")
    app_mati = app_mod.buat_app({"SECRET_KEY": "k" * 64, "ADMIN_IP_BOLEH": "103.84.116.0/24", "PRODI_BANGUN": "mati"})
    app, su = app_mati, None
    su = klien("superuji")
    teks = su.get("/").get_data(as_text=True)
    cek("Belum ada data www" in teks, "dasbor tanpa data www")
    berkas = []
    for nama in impor_www.JENIS:
        with open(os.path.join(impor_www.DATA_AWAL, nama + ".json"), "rb") as f:
            berkas.append((nama + ".json", f.read()))
    r, token = unggah(su, berkas)
    teks = su.get("/super/tarik?kiriman=%s" % token).get_data(as_text=True)
    cek("Salinan pertama: 97 prodi" in teks, "pratinjau salinan pertama")
    r = kirim(su, "/super/tarik/terapkan", sidik=isi_formulir(teks), kiriman=token)
    cek(r.status_code == 302 and len(registri.baca()["prodi"]) == 97, "unggahan pertama membuat registri 97 prodi")
    cek("97 prodi didaftarkan dari daftar slug" in su.get("/super/tarik").get_data(as_text=True), "pesan registri awal")
    cek(any(x["aksi"] == "registri-awal" and x.get("jumlah") == 97 for x in audit.baca()), "registri awal tercatat")
finally:
    shutil.rmtree(KERJA, ignore_errors=True)

for g in gagal:
    print("GAGAL", g)
print("uji_impor_www: %s" % ("LULUS" if not gagal else "%d gagal" % len(gagal)))
sys.exit(1 if gagal else 0)
