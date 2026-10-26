# -*- coding: utf-8 -*-
"""Bulk import of the lecturer links (dosen_url) of the programmes from a CSV or Excel file (langkah 5).

The file needs a column with the link (dosen_url, url, tautan, tautan_dosen, link) and a column naming the programme:
its PDDIKTI code (kode_pddikti, kode, kode_prodi) or its slug (slug). Other columns are ignored, so the template from
the admin (templat_csv) can be filled in and sent back as it is. An empty link means "no change"; a link is cleared on
the programme's own page.

Nothing is written here: rencana() returns what would change and every problem per row; the admin shows that and
applies it with registri.atur_banyak once the superadmin confirms.
"""
from __future__ import print_function

import csv
import io
import os
import re
import sys

ADMIN = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(ADMIN))

import registri  # noqa: E402

MAKS_BERKAS = 2 * 1024 * 1024
MAKS_BARIS = 2000
KOLOM_KODE = ("kode_pddikti", "kode", "kode_prodi")
KOLOM_SLUG = ("slug",)
KOLOM_URL = ("dosen_url", "url", "tautan", "tautan_dosen", "link")


def _kunci(teks):
    return re.sub(r"[\s\-]+", "_", str(teks or "").strip().lower())


def baca_tabel(nama_berkas, data):
    """Rows of the first sheet / the CSV as dicts with normalised column names. ValueError for anything else."""
    if not data:
        raise ValueError("Pilih berkas CSV atau Excel (.xlsx).")
    if len(data) > MAKS_BERKAS:
        raise ValueError("Berkas terlalu besar (maksimal 2 MB).")
    nama = (nama_berkas or "").lower()
    if data[:4] == b"PK\x03\x04":                        # .xlsx is a zip file
        baris = _baca_xlsx(data)
    elif nama.endswith((".csv", ".txt")) or not nama:
        baris = _baca_csv(data)
    elif nama.endswith(".xls"):
        raise ValueError("Format .xls lama tidak didukung; simpan ulang sebagai .xlsx atau .csv.")
    else:
        raise ValueError("Berkas harus CSV atau Excel (.xlsx).")
    baris = [[("" if v is None else str(v)).strip() for v in b] for b in baris]
    baris = [b for b in baris if any(b)]
    if not baris:
        raise ValueError("Berkas kosong.")
    if len(baris) > MAKS_BARIS + 1:
        raise ValueError("Terlalu banyak baris (maksimal %d)." % MAKS_BARIS)
    kepala = [_kunci(k) for k in baris[0]]
    return [dict(zip(kepala, b)) for b in baris[1:]]


def _baca_csv(data):
    for kode in ("utf-8-sig", "cp1252", "latin-1"):     # latin-1 never fails: the last resort
        try:
            teks = data.decode(kode)
            break
        except UnicodeDecodeError:
            continue
    if "\x00" in teks:
        raise ValueError("Berkas bukan teks CSV.")
    try:
        dialek = csv.Sniffer().sniff(teks[:4096], delimiters=",;\t")
    except csv.Error:
        dialek = csv.excel
    return list(csv.reader(io.StringIO(teks), dialek))


def _baca_xlsx(data):
    try:
        from openpyxl import load_workbook
        buku = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lembar = buku.worksheets[0]
        hasil = []
        for i, b in enumerate(lembar.iter_rows(values_only=True)):
            if i > MAKS_BARIS + 1:
                break
            hasil.append(list(b))
        buku.close()
        return hasil
    except ValueError:
        raise
    except Exception:
        raise ValueError("Berkas Excel rusak atau tidak dapat dibaca.")


def _ambil(baris, nama_kolom):
    for k in nama_kolom:
        if baris.get(k):
            return baris[k]
    return ""


def rencana(baris, daftar_registri, nama_prodi):
    """What the rows would change. Returns {"ubah": [{"kode", "nama", "lama", "baru"}], "sama": n, "kosong": n,
    "galat": [{"baris": number in the file, "pesan"}]}. nama_prodi: {code: display name}."""
    per_kode = {r["kode_pddikti"]: r for r in daftar_registri}
    per_slug = {r["slug"]: r for r in daftar_registri}
    hasil = {"ubah": [], "sama": 0, "kosong": 0, "galat": []}
    if baris and not any(k in baris[0] for k in KOLOM_URL):
        hasil["galat"].append({"baris": 1, "pesan": "Kolom tautan tidak ditemukan (dosen_url)."})
        return hasil
    if baris and not any(k in baris[0] for k in KOLOM_KODE + KOLOM_SLUG):
        hasil["galat"].append({"baris": 1, "pesan": "Kolom prodi tidak ditemukan (kode_pddikti atau slug)."})
        return hasil
    dilihat = {}
    for i, b in enumerate(baris, start=2):
        kode, slug, url = _ambil(b, KOLOM_KODE), _ambil(b, KOLOM_SLUG), _ambil(b, KOLOM_URL)
        r = per_kode.get(kode) if kode else per_slug.get(slug)
        if r is None:
            hasil["galat"].append({"baris": i, "pesan": "Prodi %s tidak terdaftar." % (kode or slug or "(kosong)")})
            continue
        if not url:
            hasil["kosong"] += 1
            continue
        masalah = registri.periksa_dosen_url(url)
        if masalah:
            hasil["galat"].append({"baris": i, "pesan": "%s: %s" % (nama_prodi.get(r["kode_pddikti"], r["kode_pddikti"]), masalah)})
            continue
        if r["kode_pddikti"] in dilihat and dilihat[r["kode_pddikti"]] != url:
            hasil["galat"].append({"baris": i, "pesan": "%s muncul lebih dari sekali dengan tautan berbeda." %
                                   nama_prodi.get(r["kode_pddikti"], r["kode_pddikti"])})
            continue
        dilihat[r["kode_pddikti"]] = url
        if (r.get("dosen_url") or "") == url:
            hasil["sama"] += 1
        else:
            hasil["ubah"].append({"kode": r["kode_pddikti"], "nama": nama_prodi.get(r["kode_pddikti"], ""),
                                  "lama": r.get("dosen_url") or "", "baru": url})
    # a row that appeared twice with the same link counts once
    hasil["ubah"] = list({u["kode"]: u for u in hasil["ubah"]}.values())
    return hasil


def _aman(teks):
    """A CSV cell that a spreadsheet will not run as a formula."""
    teks = str(teks or "")
    return "'" + teks if teks[:1] in ("=", "+", "-", "@", "\t", "\r") else teks


def templat_csv(daftar_registri, info):
    """CSV of every registered programme with its current link, for filling in. info: {code: {"jenjang", "nama",
    "unit"}}."""
    buf = io.StringIO()
    tulis = csv.writer(buf)
    tulis.writerow(["kode_pddikti", "jenjang", "nama", "fakultas", "slug", "dosen_url"])
    for r in sorted(daftar_registri, key=lambda r: r["slug"]):
        x = info.get(r["kode_pddikti"], {})
        tulis.writerow([_aman(v) for v in (r["kode_pddikti"], x.get("jenjang", ""), x.get("nama", ""), x.get("unit", ""),
                                           r["slug"], r.get("dosen_url") or "")])
    return "\ufeff" + buf.getvalue()                    # the BOM makes Excel read it as UTF-8
