# -*- coding: utf-8 -*-
"""Filling a large table of a page from a spreadsheet (Fase 2 langkah 7, K-P34): courses, partnerships, research,
publications, community service, outputs, achievements, staff, surveys, schedules and the like.

    /prodi/<kode>/halaman/<id>/impor                 the tables of a page that can be imported, with their templates;
                                                     POST: read an uploaded file and show what it holds (nothing
                                                     stored yet; the rows wait for this account, as superadmin.py's
                                                     previews do)
    /prodi/<kode>/halaman/<id>/impor/<tabel>.xlsx    the template of a table, holding the rows the page has now
    /prodi/<kode>/halaman/<id>/impor/terapkan        POST: put the previewed rows into the page's draft, after the rows
                                                     there (rows already there are not added twice) or instead of them

A table qualifies when the site shows it as a table, it may hold at least MIN_BARIS rows, and none of its required
columns is a file. File columns (PDF, photo) are not in the spreadsheet: they are attached in the editor, and a row
that comes back unchanged keeps its files. Every cell is read as the editor reads a form field and every row is checked
with the page's own checks (isi_halaman), spam terms included; rows with problems are listed and left out. The rows go
into the draft, so the page is published as always (K-P24): by the programme's admin, after a preview.
"""
from __future__ import print_function

import copy
import datetime
import io
import json
import os
import re
import sys

ADMIN = os.path.dirname(os.path.abspath(__file__))
PLATFORM = os.path.dirname(ADMIN)
sys.path.insert(0, PLATFORM)
sys.path.insert(0, ADMIN)

from flask import abort, flash, g, make_response, redirect, render_template, request, url_for  # noqa: E402

import impor_dosen  # noqa: E402
import isi_halaman as ih  # noqa: E402
import penyunting  # noqa: E402
from akses import akses  # noqa: E402

MIN_BARIS = 50                 # a table this large is tedious to fill row by row
MAKS_KOLOM = 40                # columns read from a sheet (a table has at most 8)
MEDIA = ("dokumen", "foto")
BENAR = ("ya", "y", "1", "benar", "true", "x", "v", "✓", "✔")
SALAH = ("tidak", "t", "0", "salah", "false", "-")
CONTOH_BARIS = 20              # rows of the file shown in the preview
TIPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def tabel_halaman(id_halaman):
    """[(key, node)] of the page's tables that can be filled from a spreadsheet, in the page's order."""
    templat = ih.TEMPLAT.get(id_halaman)
    hasil = []
    for kunci, node in (templat or {}).get("bidang", []):
        if node["jenis"] != "daftar" or node["isi"]["jenis"] != "objek" or node["tampil"] != "tabel" \
                or node["maks"] < MIN_BARIS:
            continue
        if any(n["jenis"] in MEDIA and n.get("wajib") for _k, n in node["isi"]["bidang"]):
            continue
        hasil.append((kunci, node))
    return hasil


def kolom(node):
    """[(key, column node)] of a table that a spreadsheet holds."""
    return [(k, n) for k, n in node["isi"]["bidang"] if n["jenis"] not in MEDIA]


def kolom_berkas(node):
    """Labels of a table's file columns (PDF, photo), which the spreadsheet leaves out."""
    return [n["label"] for _k, n in node["isi"]["bidang"] if n["jenis"] in MEDIA]


def _ribuan(x):
    return "{:,}".format(x).replace(",", ".")


def petunjuk(n):
    """What a column takes, in words (the template's instructions and the import page)."""
    j = n["jenis"]
    if j == "teks":
        return "Teks, paling panjang %s karakter" % _ribuan(n["maks"])
    if j == "bilangan":
        if n.get("format") == "tahun":
            return "Tahun, empat angka (mis. 2025)"
        return "Angka bulat %s sampai %s" % (_ribuan(n["minimum"]), _ribuan(n["maksimum"]))
    if j == "persen":
        return "Angka 0 sampai 100, paling banyak dua angka di belakang koma (mis. 87,5)"
    if j == "url":
        return "Alamat lengkap yang diawali https://"
    if j == "pilihan":
        return "Salah satu: %s" % "; ".join(n["opsi"])
    if j == "centang":
        return "ya atau tidak (kosong berarti tidak)"
    return ""


def _nama_kolom(teks):
    """A header compared without case, spaces or signs ("Mata kuliah", "mata_kuliah" and "MATA KULIAH" are one)."""
    return re.sub(r"[^0-9a-z]+", "", str(teks or "").lower())


# ------------------------------------------------------------------ reading a file

class Persen(float):
    """A number from a cell Excel shows as a percentage (0.875 shown as 87,5%), already multiplied by 100."""


def _sel_xlsx(c):
    v = getattr(c, "value", None)
    if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (getattr(c, "number_format", "") or ""):
        return Persen(round(v * 100, 6))
    return v


def _baca_xlsx(data):
    try:
        from openpyxl import load_workbook
        buku = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lembar = buku.worksheets[0]
        hasil = [[_sel_xlsx(c) for c in b] for b in lembar.iter_rows(max_row=impor_dosen.MAKS_BARIS + 2,
                                                                       max_col=MAKS_KOLOM)]
        buku.close()
        return hasil
    except Exception:
        raise ValueError("Berkas Excel rusak atau tidak dapat dibaca.")


def teks_sel(v):
    """A cell as text: whole numbers without ".0", a date or a time Excel recognised in a text column as dd/mm/yyyy
    or hh:mm."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "ya" if v else "tidak"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y %H:%M") if (v.hour, v.minute) != (0, 0) else v.strftime("%d/%m/%Y")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    return str(v).strip()


def baca_berkas(nama_berkas, data):
    """[(row number in the file, [cell, ...])] of the first sheet of an Excel file or of a CSV, empty rows left out;
    the header is the first. ValueError with a message for the user for any other file."""
    if not data:
        raise ValueError("Pilih berkas Excel (.xlsx) atau CSV.")
    if len(data) > impor_dosen.MAKS_BERKAS:
        raise ValueError("Berkas terlalu besar (maksimal 2 MB).")
    nama = (nama_berkas or "").lower()
    if data[:4] == b"PK\x03\x04":                        # .xlsx is a zip file
        baris = _baca_xlsx(data)
    elif nama.endswith((".csv", ".txt")):
        baris = impor_dosen._baca_csv(data)
    elif nama.endswith(".xls"):
        raise ValueError("Format .xls lama tidak didukung; simpan ulang sebagai .xlsx.")
    else:
        raise ValueError("Berkas harus Excel (.xlsx) atau CSV.")
    hasil = [(i, b[:MAKS_KOLOM]) for i, b in enumerate(baris, 1) if any(teks_sel(v) for v in b)]
    if not hasil:
        raise ValueError("Berkas kosong.")
    if len(hasil) > impor_dosen.MAKS_BARIS + 1:
        raise ValueError("Terlalu banyak baris (maksimal %d)." % impor_dosen.MAKS_BARIS)
    return hasil


def nilai_sel(n, v):
    """(value, problem) of a cell for a column: the value as the editor stores it, or what is wrong with it in words.
    An empty cell is (None, None)."""
    j = n["jenis"]
    if j == "centang":
        s = teks_sel(v).lower()
        if s in BENAR:
            return True, None
        return (None, None) if s in SALAH or not s else (None, "isi ya atau tidak")
    if j in ("bilangan", "persen") and isinstance(v, (int, float)) and not isinstance(v, bool):
        if j == "persen":
            x = float(v)
            return (int(x) if x.is_integer() else x), None
        return (int(v), None) if float(v).is_integer() else (None, "harus angka bulat")
    s = teks_sel(v)
    if j in ("teks", "url"):
        s = penyunting._teks(s, j == "url" or not n["panjang"])
        return (s or None), None
    s = s.strip()
    if not s:
        return None, None
    if j == "pilihan":
        cocok = {" ".join(o.lower().split()): o for o in n["opsi"]}
        o = cocok.get(" ".join(s.lower().split()))
        return (o, None) if o else (None, "harus salah satu dari: %s" % "; ".join(n["opsi"]))
    if j == "bilangan":
        if penyunting.RIBUAN.match(s):
            s = s.replace(".", "")
        return (int(s), None) if penyunting.BULAT.match(s) else (None, "harus angka bulat")
    if j == "persen":
        s = s.rstrip("%").strip()
        if penyunting.PECAHAN.match(s):
            x = float(s.replace(",", "."))
            return (int(x) if x.is_integer() else x), None
        return None, "harus angka 0 sampai 100"
    return None, "jenis kolom tidak dikenal"


def rencana(node, baris, kode):
    """What a file read with baca_berkas() holds for a table: {"baris": [row, ...] that pass the checks, "galat":
    ["Baris 7, Mata kuliah: wajib diisi", ...], "bermasalah": rows left out, "dibaca": rows read, "kolom": the
    table's columns found, "abaikan": headers that are not the table's}. ValueError when the file cannot be used at
    all: no header the table knows, a required column missing, more rows than the table may hold."""
    kol = kolom(node)
    jenis = dict(kol)
    nama = {}
    for k, n in kol:
        nama.setdefault(_nama_kolom(k), k)
    for k, n in kol:
        nama[_nama_kolom(n["label"])] = k                 # a label wins over another column's key
    _no, kepala = baris[0]
    peta, abaikan = {}, []
    for i, h in enumerate(kepala):
        h = teks_sel(h)
        k = nama.get(_nama_kolom(h))
        if not h:
            continue
        if k is None or k in peta.values():
            abaikan.append(h)
        else:
            peta[i] = k
    if not peta:
        raise ValueError("baris pertama berkas harus berisi nama kolom tabel ini (%s). Pakai templat yang bisa diunduh "
                         "di halaman ini." % ", ".join(n["label"] for _k, n in kol))
    hilang = [n["label"] for k, n in kol if n.get("wajib") and k not in peta.values()]
    if hilang:
        raise ValueError("kolom wajib tidak ada: %s." % ", ".join(hilang))
    isi = baris[1:]
    if len(isi) > node["maks"]:
        raise ValueError("berkas berisi %d baris, padahal tabel ini paling banyak %d baris." % (len(isi), node["maks"]))
    hasil = {"baris": [], "galat": [], "bermasalah": 0, "dibaca": 0, "abaikan": abaikan,
             "kolom": [jenis[k]["label"] for _i, k in sorted(peta.items())]}
    for no, sel in isi:
        r, salah = {}, {}
        for i, k in peta.items():
            v, m = nilai_sel(jenis[k], sel[i] if i < len(sel) else None)
            if m:
                salah[jenis[k]["label"]] = m
            elif v is not None:
                r[k] = v
        if not r and not salah:
            continue                                     # something only in columns that are not the table's
        hasil["dibaca"] += 1
        tempat = "Baris %d" % no
        pesan = ["%s, %s: %s" % (tempat, label, m) for label, m in salah.items()]
        pesan += [p for p in ih.masalah_nilai(node["isi"], r, tempat, kode)
                  if not any(p.startswith("%s, %s:" % (tempat, label)) for label in salah)]
        if pesan:
            hasil["galat"] += pesan
            hasil["bermasalah"] += 1
        else:
            hasil["baris"].append(r)
    return hasil


# ------------------------------------------------------------------ the rows of the table

def _kunci_baris(node, r):
    """A row by the values of its spreadsheet columns ("", nothing and "no" alike; spaces as typed do not count)."""
    def nilai(v):
        if v in ("", None, False):
            return None
        return " ".join(v.split()) if isinstance(v, str) else v
    return json.dumps([nilai(r.get(k)) for k, _n in kolom(node)], sort_keys=True)


def ringkasan(node, lama, baru):
    """For the preview: the rows of the file the table already holds, its size after each choice, and the files
    (PDF, photo) that replacing the rows would leave out of the draft."""
    lama = [r for r in (lama or []) if isinstance(r, dict)]
    ada = {_kunci_baris(node, r) for r in lama}
    di_berkas = {_kunci_baris(node, r) for r in baru}
    media = [k for k, n in node["isi"]["bidang"] if n["jenis"] in MEDIA]
    sama = sum(1 for r in baru if _kunci_baris(node, r) in ada)
    return {"lama": len(lama), "sama": sama, "tambah": len(lama) + len(baru) - sama, "ganti": len(baru),
            "lepas": sum(1 for r in lama for k in media if r.get(k) and _kunci_baris(node, r) not in di_berkas),
            "maks": node["maks"]}


def gabung(node, lama, baru, mode):
    """The table after an import. "tambah": the new rows after the ones there, a row already there not again;
    "ganti": the new rows only, a row that was there keeping its files. ValueError beyond the table's size."""
    lama = [r for r in (lama or []) if isinstance(r, dict)]
    media = [k for k, n in node["isi"]["bidang"] if n["jenis"] in MEDIA]
    per_kunci = {}
    for r in lama:
        per_kunci.setdefault(_kunci_baris(node, r), r)
    if mode == "tambah":
        hasil = lama + [r for r in baru if _kunci_baris(node, r) not in per_kunci]
    elif mode == "ganti":
        hasil = []
        for r in baru:
            asal = per_kunci.get(_kunci_baris(node, r)) or {}
            hasil.append(dict(r, **{k: copy.deepcopy(asal[k]) for k in media if asal.get(k)}))
    else:
        raise ValueError("pilihan impor tidak dikenal")
    if len(hasil) > node["maks"]:
        raise ValueError("tabel ini paling banyak %d baris, sedangkan hasilnya %d baris" % (node["maks"], len(hasil)))
    return hasil


# ------------------------------------------------------------------ the template

def templat_xlsx(node, baris_ada, judul):
    """The spreadsheet to fill: the "Data" sheet (header, the rows the page holds now, drop-down lists for choices and
    ya/tidak), a "Petunjuk" sheet, the choices on a hidden sheet. A text is always written as text: one that starts
    with "=" never becomes a formula."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter, quote_sheetname
    from openpyxl.worksheet.datavalidation import DataValidation

    kol = kolom(node)
    buku = Workbook()
    data = buku.active
    data.title = "Data"
    data.append([n["label"] for _k, n in kol])
    for c in data[1]:
        c.font = Font(bold=True)
    data.freeze_panes = "A2"
    for r in baris_ada:
        if not isinstance(r, dict):
            continue
        sel = []
        for k, n in kol:
            v = r.get(k)
            if n["jenis"] == "centang":
                v = "ya" if v is True else None
            elif isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            elif isinstance(v, bool) or not isinstance(v, (int, float)):
                v = None
            sel.append(v)
        data.append(sel)
        for c in data[data.max_row]:
            if isinstance(c.value, str):
                c.data_type = "s"
    pilihan = buku.create_sheet("Pilihan")
    pilihan.sheet_state = "hidden"
    akhir = node["maks"] + 1
    for i, (k, n) in enumerate(kol, 1):
        huruf = get_column_letter(i)
        data.column_dimensions[huruf].width = 40 if n["jenis"] in ("teks", "url") and n.get("maks", 300) > 60 else 16
        if n["jenis"] not in ("pilihan", "centang"):
            continue
        opsi = n["opsi"] if n["jenis"] == "pilihan" else ["ya", "tidak"]
        for j, o in enumerate(opsi, 1):
            pilihan.cell(row=j, column=i, value=o)
        dv = DataValidation(type="list", allow_blank=True, showErrorMessage=True, errorTitle=n["label"],
                            error="Pilih salah satu isi dari daftar.",
                            formula1="%s!$%s$1:$%s$%d" % (quote_sheetname("Pilihan"), huruf, huruf, len(opsi)))
        dv.add("%s2:%s%d" % (huruf, huruf, akhir))
        data.add_data_validation(dv)
    pt = buku.create_sheet("Petunjuk", 1)
    pt.column_dimensions["A"].width = 28
    pt.column_dimensions["B"].width = 10
    pt.column_dimensions["C"].width = 80
    berkas = kolom_berkas(node)
    for baris in (["Petunjuk: %s" % judul], [],
                  ["Isi lembar Data, satu baris tabel per baris Excel, paling banyak %d baris." % node["maks"]],
                  ["Baris pertama (nama kolom) jangan diubah. Baris yang sudah ada di halaman ikut tertulis di bawahnya."],
                  ["Kolom %s tidak ada di Excel: unggah berkasnya lewat formulir halaman sesudah impor." % ", ".join(berkas)]
                  if berkas else None,
                  ["Hasil impor masuk ke draf halaman. Periksa di penyunting, lalu ajukan atau terbitkan seperti biasa."],
                  [], ["Kolom", "Wajib", "Isi"]):
        if baris is not None:
            pt.append(baris)
    pt.cell(row=1, column=1).font = Font(bold=True)
    for c in pt[pt.max_row]:
        c.font = Font(bold=True)
    for _k, n in kol:
        pt.append([n["label"], "ya" if n.get("wajib") else "", petunjuk(n)])
        pt.cell(row=pt.max_row, column=3).alignment = Alignment(wrap_text=True)
    keluar = io.BytesIO()
    buku.save(keluar)
    return keluar.getvalue()


# ------------------------------------------------------------------ the admin

def pasang(app, alat):

    def cari(kode, id_, tabel=None):
        """(the page's menu entry, its importable tables {key: node}); 404 for a page without one."""
        tabel_ada = dict(tabel_halaman(id_))
        if not tabel_ada:
            abort(404)
        ini = alat.ini_halaman(kode, id_)
        if tabel is not None and tabel not in tabel_ada:
            abort(404)
        return ini, tabel_ada

    def baris_sekarang(kode, id_, tabel):
        return [r for r in ih.isi_kerja(ih.baca(kode, id_)).get(tabel) or [] if isinstance(r, dict)]

    def tampil(kode, id_, kode_status=200):
        ini, tabel_ada = cari(kode, id_)
        daftar = [{"kunci": k, "node": n, "kolom": kolom(n), "berkas": kolom_berkas(n),
                   "jumlah": len(baris_sekarang(kode, id_, k))} for k, n in tabel_ada.items()]
        return render_template("halaman_impor.html", pr=g.prodi, ini=ini, daftar=daftar, petunjuk=petunjuk,
                               maks_baris=impor_dosen.MAKS_BARIS), kode_status

    @app.route("/prodi/<kode>/halaman/<id_>/impor")
    @akses("prodi")
    def isi_impor(kode, id_):
        return tampil(kode, id_)

    @app.route("/prodi/<kode>/halaman/<id_>/impor/<tabel>.xlsx")
    @akses("prodi")
    def isi_impor_templat(kode, id_, tabel):
        ini, tabel_ada = cari(kode, id_, tabel)
        node = tabel_ada[tabel]
        isi = templat_xlsx(node, baris_sekarang(kode, id_, tabel), "%s — halaman %s, %s" % (
            node["label"], ini["label"], g.prodi["nama"]))
        resp = make_response(isi)
        resp.headers["Content-Type"] = TIPE_XLSX
        resp.headers["Content-Disposition"] = 'attachment; filename="%s-%s-%s.xlsx"' % (kode, id_, tabel)
        return resp

    @app.route("/prodi/<kode>/halaman/<id_>/impor", methods=["POST"])
    @akses("prodi")
    def isi_impor_baca(kode, id_):
        ini, tabel_ada = cari(kode, id_)
        tabel = request.form.get("tabel", "")
        if tabel not in tabel_ada:
            abort(400)
        node = tabel_ada[tabel]
        berkas = request.files.get("berkas")
        try:
            if not berkas or not berkas.filename:
                raise ValueError("pilih berkas Excel (.xlsx) atau CSV.")
            hasil = rencana(node, baca_berkas(berkas.filename, berkas.read(impor_dosen.MAKS_BERKAS + 1)), kode)
        except ValueError as e:
            flash("%s: berkas tidak dapat diimpor: %s" % (node["label"], e), "galat")
            return tampil(kode, id_, 400)
        lama = baris_sekarang(kode, id_, tabel)
        token = alat.simpan_sementara("impor-tabel", {"kode": kode, "halaman": id_, "tabel": tabel,
                                                      "baris": hasil["baris"]}) if hasil["baris"] else None
        alat.catat("isi-impor-pratinjau", prodi=kode, halaman=id_, tabel=tabel, berkas=berkas.filename[:100],
                   baris=len(hasil["baris"]), bermasalah=hasil["bermasalah"])
        return render_template("halaman_impor_pratinjau.html", pr=g.prodi, ini=ini, node=node, tabel=tabel,
                               hasil=hasil, r=ringkasan(node, lama, hasil["baris"]), token=token,
                               nama_berkas=berkas.filename, kolom=kolom(node), contoh=hasil["baris"][:CONTOH_BARIS],
                               maks_galat=100, berkas_kolom=kolom_berkas(node))

    @app.route("/prodi/<kode>/halaman/<id_>/impor/terapkan", methods=["POST"])
    @akses("prodi")
    def isi_impor_terapkan(kode, id_):
        _ini, tabel_ada = cari(kode, id_)
        ke_impor = url_for("isi_impor", kode=kode, id_=id_)
        simpanan = alat.ambil_sementara("impor-tabel", request.form.get("token"))
        if not simpanan or simpanan.get("kode") != kode or simpanan.get("halaman") != id_ \
                or simpanan.get("tabel") not in tabel_ada:
            flash("Pratinjau impor sudah tidak berlaku (dipakai, kedaluwarsa, atau milik akun lain). Unggah ulang "
                  "berkasnya.", "galat")
            return redirect(ke_impor)
        tabel, mode, baru = simpanan["tabel"], request.form.get("mode", ""), simpanan["baris"]
        node = tabel_ada[tabel]
        data = ih.baca(kode, id_)
        isi = copy.deepcopy(ih.isi_kerja(data))
        lama = [r for r in isi.get(tabel) or [] if isinstance(r, dict)]
        try:
            isi[tabel] = gabung(node, lama, baru, mode)
            ih.simpan_draf(kode, id_, isi, g.akun["nama"], data.get("versi", 0))
        except (ValueError, ih.GagalIsi) as e:
            flash("Impor dibatalkan: %s. Unggah ulang berkasnya." % e, "galat")
            return redirect(ke_impor)
        alat.catat("isi-tabel-diimpor", prodi=kode, halaman=id_, tabel=tabel, mode=mode, baris=len(baru),
                   jumlah=len(isi[tabel]))
        if mode == "tambah":
            masuk = len(isi[tabel]) - len(lama)
            pesan = "%d baris ditambahkan" % masuk + (
                " (%d baris yang sudah ada dilewati)" % (len(baru) - masuk) if len(baru) > masuk else "")
        else:
            pesan = "tabel diganti dengan %d baris dari berkas" % len(baru)
        flash("%s: %s; tabel kini %d baris. Ini masih draf: situs belum berubah sampai halaman diterbitkan seperti "
              "biasa." % (node["label"], pesan, len(isi[tabel])), "ok")
        return redirect(url_for("isi_sunting", kode=kode, id_=id_))
